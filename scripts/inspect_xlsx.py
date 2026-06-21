import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('input/course_spec.xlsx', data_only=True)
print('Sheets:', wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f'=== Sheet: {sn} | max_row={ws.max_row} max_col={ws.max_column} ===')

# Dump Sheet "专业实验" (index 2)
ws = wb[wb.sheetnames[2]]
print(f'\n--- Dumping {wb.sheetnames[2]} ---')
for r in range(1, min(ws.max_row + 1, 70)):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v is not None:
            s = str(v).replace('\n', '\\n')
            print(f'[{r},{c}] {s[:300]}')
    print('---')
